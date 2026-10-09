"""Adds the 'Verification 08-10-2026' sheet: Workload Report values observed after import vs Monthly Expected."""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

F = sys.argv[1]
wb = openpyxl.load_workbook(F)
FONT = 'Arial'
vs = wb.create_sheet('Verification 08-10-2026', 1)

# Observed in Reports > Workload Report > All Users, Allocated Man Days, Bar, month periods (08-10-2026).
OBS = {
    'Oct 2026': {'Unassigned': 2.222, 'UIAutomation_Workload_Resource': 15.6, 'Vishal customprofile': 15.333,
                 'Vishal liteuser': 12.25, 'vishal PL manager': 38.667, 'Vishal PLAdmin': 56.583,
                 'Vishal plemployee': 15, 'Vishal plreadonly': 13.333, 'Vishal workload': 131.333},
    'Nov 2026': {'Unassigned': 0, 'UIAutomation_Workload_Resource': 14.333, 'Vishal customprofile': 11.25,
                 'Vishal liteuser': 13.333, 'vishal PL manager': 11.75, 'Vishal PLAdmin': 13.333,
                 'Vishal plemployee': 23.833, 'Vishal plreadonly': 16, 'Vishal workload': 20},
    'Dec 2026': {'Unassigned': 0, 'UIAutomation_Workload_Resource': 14.25, 'Vishal customprofile': 14.667,
                 'Vishal liteuser': 16.722, 'vishal PL manager': 19.333, 'Vishal PLAdmin': 15.056,
                 'Vishal plemployee': 10.25, 'Vishal plreadonly': 16.333, 'Vishal workload': 14.333},
}
WHY = {
    ('Oct 2026', 'Vishal workload'): 'Other projects (NI-6, NI-1) = +549 h (pre-import total 1044 h - previous import 495 h).',
    ('Oct 2026', 'Vishal PLAdmin'): '+198 h other projects (pre-import 396 h - previous import 198 h); -22.5 h FINDING F1 (150 %/day capped to 100 %, Oct W2).',
    ('Oct 2026', 'vishal PL manager'): '+121 h UI1-T1 / other projects (pre-import 220 h - previous import 99 h); +8 h FINDING F2 (WL-E03 weekend-only task allocates 4 h on Sat 3 + Sun 4 Oct).',
    ('Oct 2026', 'Vishal liteuser'): '-22.5 h FINDING F1 (150 %/day capped to 100 %, Oct over week).',
    ('Oct 2026', 'Vishal customprofile'): '-9 h FINDING F1 (WL-E05 150 %/day stored as 18:00 = 100 %).',
    ('Nov 2026', 'Vishal customprofile'): '-22.5 h FINDING F1 (150 %/day capped, Nov over week).',
    ('Nov 2026', 'vishal PL manager'): '-22.5 h FINDING F1 (150 %/day capped, Nov over week).',
    ('Nov 2026', 'Vishal plemployee'): '+45 h: two NI-6 tasks at 4:30/day on 2-6 Nov (expanded row, verified).',
    ('Nov 2026', 'Vishal workload'): '+22.5 h: other project tasks (not expanded - verify).',
    ('Dec 2026', 'UIAutomation_Workload_Resource'): '-18 h FINDING F1 (150 %/day capped, Dec 28-31 over week, 4 days).',
    ('Dec 2026', 'Vishal plemployee'): '-18 h FINDING F1 (150 %/day capped, Dec 28-31 over week, 4 days).',
    ('Dec 2026', 'vishal PL manager'): '+22.5 h: NI-6 task 4:30/day on 7-11 Dec (expanded row, verified).',
}
vs.append(['Workload Report vs expected - observed 08-10-2026 (All Users, Allocated Man Days, month periods)'])
vs.append(['Expected = Monthly Expected (this project, previous import included). Report also contains other projects and UI1-T1, explained in the last column.'])
vs.append([])
vs.append(['Month', 'User', 'Report (man-days)', 'Report (hrs)', 'Expected (hrs)', 'Delta (hrs)', 'Result', 'Explanation'])
for c in vs[4]:
    c.font = Font(name=FONT, bold=True)
for month, users in OBS.items():
    for u, d in users.items():
        r = vs.max_row + 1
        vs.append([month, u, d])
        vs[f'D{r}'] = f"=C{r}*'Assignment Calc'!$B$1"
        vs[f'E{r}'] = (f"=SUMPRODUCT(('Monthly Expected'!$A$2:$A$100=B{r})*('Monthly Expected'!$B$2:$B$100=A{r})"
                       f"*'Monthly Expected'!$F$2:$F$100)")
        vs[f'F{r}'] = f'=ROUND(D{r}-E{r},1)'
        why = WHY.get((month, u), '')
        vs[f'G{r}'] = f'=IF(ABS(F{r})<0.2,"Match",IF(ISNUMBER(SEARCH("FINDING",H{r})),"Explained - finding","Explained - other data"))'
        vs[f'H{r}'] = why
        for c in vs[r]:
            c.font = Font(name=FONT)
        vs[f'C{r}'].font = Font(name=FONT, color='0000FF')
        for col in 'DEF':
            vs[f'{col}{r}'].number_format = '0.00'
last = vs.max_row
vs.append([])
vs.append(['FINDINGS'])
vs[f'A{vs.max_row}'].font = Font(name=FONT, bold=True)
findings = [
    ('F1', 'Import caps Work %/day at 100 %. WL-E05 (150 %/day, 2 days) was stored as Work Hours 18:00 instead of 27:00; every 150 %/day row lost 4.5 h/day. No warning in the import wizard.'),
    ('F2', 'A task whose whole range is non-working days (WL-E03, Sat 3 - Sun 4 Oct, 4:00 hrs/day) is allocated 4 h on Saturday and Sunday (0.444d each) although business hours are Mon-Fri.'),
    ('F3', 'Associated Team and Tags columns did not import (checked UI1-T114 and UI1-T15: both empty; file had UIAutomation_Workload_TeamA/B and wl-edge / wl-over). Custom fields, priority, billing, % completed, dates and work hours did import.'),
    ('F4', 'Import creates a new task list even when one with the same name exists (second "Workload CV List" and second "General").'),
    ('F5', 'Importer reads rows from EVERY sheet of the workbook (Scenario Matrix text appeared as statuses) - keep the import file to a single Tasks sheet.'),
    ('OK', 'Unassigned task with hours (WL-E10) shows 4:00/day under Unassigned User; resource, read-only and lite users can own tasks; multi-owner hrs and % split correctly (WL-E08, WL-E09); cross-month and weekend-span totals split correctly (WL-E06, WL-E01).'),
]
for k, t in findings:
    vs.append([k, t])
    vs[f'A{vs.max_row}'].font = Font(name=FONT, bold=True)
    vs[f'B{vs.max_row}'].font = Font(name=FONT)
    vs.merge_cells(f'B{vs.max_row}:H{vs.max_row}')
    vs[f'B{vs.max_row}'].alignment = Alignment(wrap_text=True, vertical='top')
    vs.row_dimensions[vs.max_row].height = 32
from openpyxl.formatting.rule import FormulaRule
vs.conditional_formatting.add(f'G5:G{last}', FormulaRule(formula=['G5="Match"'], fill=PatternFill('solid', fgColor='B7E1B5')))
vs.conditional_formatting.add(f'G5:G{last}', FormulaRule(formula=['ISNUMBER(SEARCH("finding",G5))'], fill=PatternFill('solid', fgColor='F4B6B6')))
vs.conditional_formatting.add(f'G5:G{last}', FormulaRule(formula=['ISNUMBER(SEARCH("other",G5))'], fill=PatternFill('solid', fgColor='FFE29A')))
for col, w in zip('ABCDEFGH', [10, 32, 12, 12, 13, 11, 22, 110]):
    vs.column_dimensions[col].width = w
vs['A1'].font = Font(name=FONT, bold=True, size=13)
vs['A2'].font = Font(name=FONT, italic=True)
vs.freeze_panes = 'C5'
wb.save(F)
print('ok')
