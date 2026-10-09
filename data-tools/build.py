"""Builds the full-coverage Workload import workbook (Q4 2026, portal customworkload).

Sheet 1 'Tasks' is the Zoho Projects import sheet (same 38 columns as the original file, all text).
Other sheets compute the expected Workload Report values with live formulas.
"""
import datetime as dt
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

SRC = sys.argv[1]   # original import file (for headers, custom-field sets, previous-import baseline)
OUT = sys.argv[2]          # expected-results workbook (Tasks + calc sheets)
OUT_IMPORT = sys.argv[3]   # Tasks-only import file

AVAIL = 9.0
D = dt.date
FONT = 'Arial'

USERS = [('VW', 'Vishal workload'), ('PA', 'Vishal PLAdmin'), ('PM', 'vishal PL manager'),
         ('PE', 'Vishal plemployee'), ('RO', 'Vishal plreadonly'), ('LU', 'Vishal liteuser'),
         ('CP', 'Vishal customprofile'), ('RS', 'UIAutomation_Workload_Resource')]
UNAME = dict(USERS)
REPORT_ROWS = [n for _, n in USERS] + ['Unassigned']

MONTHS = [('Oct', 10, [(5, 9), (12, 16), (19, 23), (26, 30)]),
          ('Nov', 11, [(2, 6), (9, 13), (16, 20), (23, 27)]),
          ('Dec', 12, [(7, 11), (14, 18), (21, 25), (28, 31)])]
STATES = ['Over-allocated', 'Optimally allocated', 'Under-allocated', 'Unallocated']

# (unit, value-kind, value). kind: 'pd' hrs/day value, 'pct' %/day, 'tot' total-hours-per-working-day (x working days)
VARIANTS = {
    'Over-allocated': [
        [('hrs/day', 6.0), ('hrs/day', 4.5)],           # 10.5 h = 116.67 %
        [('%/day', 150)],                               # 13.5 h = 150 %
        [('hrs', 12.0)],                                # 12 h/day spread from total = 133.33 %
        [('hrs/day', 5.0), ('%/day', 50)],              # 9.5 h = 105.56 % (just over)
    ],
    'Optimally allocated': [
        [('hrs/day', 9.0)],
        [('%/day', 100)],
        [('hrs', 4.5), ('hrs/day', 4.5)],               # total + per-day mix = 9 h
        [('hrs', 9.0)],
    ],
    'Under-allocated': [
        [('hrs/day', 4.5)],                             # 50 %
        [('%/day', 25)],                                # 2.25 h
        [('hrs', 3.0)],                                 # 3 h/day
        [('hrs/day', 8.5)],                             # 94.44 % (just under)
    ],
    'Unallocated': [[], [], [], [('hrs/day', 0.0)]],   # variant d: owned task with 0 h
}
VARIANT_TEXT = {'Over-allocated': ['6:00/d + 4:30/d', '150 %/d', 'total 12 h/d', '5:00/d + 50 %/d'],
                'Optimally allocated': ['9:00/d', '100 %/d', 'total 4.5 h/d + 4:30/d', 'total 9 h/d'],
                'Under-allocated': ['4:30/d', '25 %/d', 'total 3 h/d', '8:30/d'],
                'Unallocated': ['no task', 'no task', 'no task', 'task with 0:00']}

STATUSES = ['Open', 'In Progress', 'In Review', 'To be Tested', 'On Hold', 'Delayed', 'Closed', 'Cancelled']
PRIORITIES = ['High', 'Medium', 'Low', 'None']
BILLING = ['Billable', 'Non Billable', 'None']
TEAMS = ['UIAutomation_Workload_TeamA', 'UIAutomation_Workload_TeamB',
         'UIAutomation_Workload_TeamB, UIAutomation_Workload_TeamA', None]
PCTS = [0, 10, 25, 50, 75, 90]
TAG = {'Over-allocated': 'wl-over', 'Optimally allocated': 'wl-optimal', 'Under-allocated': 'wl-under',
       'Unallocated': 'wl-unallocated', 'Edge': 'wl-edge'}


def hm(h):
    m = round(h * 60)
    return f'{m // 60:02d}:{m % 60:02d}'


def ds(d):
    return d.strftime('%d/%m/%Y') if d else None


def workdays(a, b):
    if not a:
        return 0
    n, d = 0, a
    while d <= b:
        n += d.weekday() < 5
        d += dt.timedelta(days=1)
    return n


# ---------------------------------------------------------------- task plan
tasks = []   # dict(name, owners=[(user, unit, value)], unit, start, end, state, month, week, tasklist, scenario, expect)


def add(name, owners, unit, start, end, state, tasklist, scenario, expect, month='', week=''):
    tasks.append(dict(name=name, owners=owners, unit=unit, start=start, end=end, state=state,
                      tasklist=tasklist, scenario=scenario, expect=expect, month=month, week=week))


plan = {}  # (user, month, week) -> (state, variant)
for o, (code, uname) in enumerate(USERS):
    for m, (mon, mm, weeks) in enumerate(MONTHS):
        v = (o + m * 3) % 4
        for w, (a, b) in enumerate(weeks):
            state = STATES[(w + o + m) % 4]
            plan[(uname, mon, w)] = (state, v)
            parts = VARIANTS[state][v]
            start, end = D(2026, mm, a), D(2026, mm, b)
            wd = workdays(start, end)
            perday = sum(val if u == 'hrs/day' else (val / 100 * AVAIL if u == '%/day' else val) for u, val in parts)
            for i, (u, val) in enumerate(parts):
                stored = val * wd if u == 'hrs' else val       # 'hrs' rows store the TOTAL
                tl = 'Workload CV List' if state == 'Unallocated' else f'WL {mon} 2026'
                n = f' ({i + 1} of {len(parts)})' if len(parts) > 1 else ''
                add(f'WL {mon} W{w + 1} {state}{n} - {uname}', [(uname, u, stored)], u, start, end, state, tl,
                    f'{state} week, variant {"abcd"[v]}: {VARIANT_TEXT[state][v]}',
                    f'{uname}: {hm(perday)} h/day on {wd} working days ({perday / AVAIL:.2%} of {AVAIL:g} h)',
                    mon, f'W{w + 1}')

E = 'WL Edge Cases'
add('WL-E01 Weekend span total hours - Vishal workload', [('Vishal workload', 'hrs', 18.0)], 'hrs',
    D(2026, 10, 1), D(2026, 10, 4), 'Edge', E, 'Thu-Sun span, 18 h total spread over working days only',
    '9:00 on Thu 1 Oct and Fri 2 Oct; 0 on Sat/Sun')
add('WL-E02 Single-day total hours - Vishal PLAdmin', [('Vishal PLAdmin', 'hrs', 12.0)], 'hrs',
    D(2026, 10, 1), D(2026, 10, 1), 'Edge', E, 'Single day task, 12 h total', '12:00 on 1 Oct = 133.33 % (over)')
add('WL-E03 Weekend-only task - vishal PL manager', [('vishal PL manager', 'hrs/day', 4.0)], 'hrs/day',
    D(2026, 10, 3), D(2026, 10, 4), 'Edge', E, 'Starts and ends on non-working days (Sat-Sun)',
    'VERIFY: 0 h expected (no working days in range)')
add('WL-E04 Zero-hour owned task - Vishal plemployee', [('Vishal plemployee', 'hrs', 0.0)], 'hrs',
    D(2026, 10, 1), D(2026, 10, 2), 'Edge', E, 'Owner set, work 00:00', 'Task listed under user, 0 h allocated')
add('WL-E05 150 percent per day - Vishal customprofile', [('Vishal customprofile', '%/day', 150)], '%/day',
    D(2026, 10, 1), D(2026, 10, 2), 'Edge', E, '%/day above 100', '13:30 per day on 1-2 Oct (over)')
add('WL-E06 Cross-month total hours - Vishal plreadonly', [('Vishal plreadonly', 'hrs', 45.0)], 'hrs',
    D(2026, 11, 30), D(2026, 12, 4), 'Edge', E, 'Mon 30 Nov - Fri 4 Dec, 45 h total',
    '9:00/day; Nov gets 9 h, Dec gets 36 h')
add('WL-E07 Resource cross-month - UIAutomation_Workload_Resource',
    [('UIAutomation_Workload_Resource', 'hrs/day', 9.0)], 'hrs/day',
    D(2026, 11, 30), D(2026, 12, 4), 'Edge', E, 'Resource as owner, crosses month end',
    '9:00/day 30 Nov - 4 Dec (optimal)')
add('WL-E08 Multi-owner different hours - Vishal liteuser and Vishal customprofile',
    [('Vishal liteuser', 'hrs/day', 6.0), ('Vishal customprofile', 'hrs/day', 3.0)], 'hrs/day',
    D(2026, 12, 1), D(2026, 12, 4), 'Edge', E, 'Two owners, different hrs/day',
    'liteuser 6:00/day, customprofile 3:00/day on 1-4 Dec')
add('WL-E09 Multi-owner percent - Vishal workload, Vishal PLAdmin, vishal PL manager',
    [('Vishal workload', '%/day', 50), ('Vishal PLAdmin', '%/day', 25), ('vishal PL manager', '%/day', 25)],
    '%/day', D(2026, 12, 1), D(2026, 12, 4), 'Edge', E, 'Three owners, different %/day',
    'workload 4:30, PLAdmin 2:15, PL manager 2:15 per day on 1-4 Dec')
add('WL-E10 Unassigned task with hours', [('Unassigned', 'hrs/day', 4.0)], 'hrs/day',
    D(2026, 10, 5), D(2026, 10, 9), 'Edge', E, 'No owner, 4:00/day',
    'VERIFY: Unassigned User row shows 4:00/day 5-9 Oct (or 0 if hours are dropped without owner)')
add('WL-E11 Unassigned task without hours', [('Unassigned', 'hrs', 0.0)], 'hrs',
    D(2026, 11, 2), D(2026, 11, 6), 'Edge', E, 'No owner, no work', 'Unassigned User task count +1, 0 h')
add('WL-E12 No dates task - Vishal workload', [('Vishal workload', 'hrs', 10.0)], 'hrs', None, None, 'Edge', E,
    'Owner and 10 h total but no start/due date', 'Not shown on any date in the Workload Report')
add('WL-E13 Single Friday full day - Vishal liteuser', [('Vishal liteuser', 'hrs/day', 9.0)], 'hrs/day',
    D(2026, 10, 2), D(2026, 10, 2), 'Edge', E, 'Single working day', '9:00 on Fri 2 Oct (optimal)')
add('WL-E14 60 percent single day - UIAutomation_Workload_Resource',
    [('UIAutomation_Workload_Resource', '%/day', 60)], '%/day',
    D(2026, 10, 1), D(2026, 10, 1), 'Edge', E, 'Resource, %/day below 100', '5:24 on 1 Oct (under)')
add('WL-E15 Multi-owner same hours - vishal PL manager and Vishal plemployee',
    [('vishal PL manager', 'hrs/day', 4.5), ('Vishal plemployee', 'hrs/day', 4.5)], 'hrs/day',
    D(2026, 11, 30), D(2026, 11, 30), 'Edge', E, 'Two owners, same hrs/day, single day',
    '4:30 each on Mon 30 Nov (under)')

# ---------------------------------------------------------------- workbook
src = openpyxl.load_workbook(SRC)
sws = src['Tasks']
srows = list(sws.values)
HDR = list(srows[0])
CUSTOM = HDR[18:]                          # Multi Select ... Single-Line
CSETS = [None] + [list(srows[i][18:]) for i in (1, 2, 3)]   # set 0 = blank

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Tasks'
ws.append(HDR)
bold = Font(name=FONT, bold=True)
for c in ws[1]:
    c.font = bold
for col, w in zip('ABCDEFGHIJKLMNOPQR', [62, 44, 14, 40, 14, 12, 12, 9, 12, 9, 11, 70, 11, 16, 13, 60, 13, 18]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'B2'

task_row = []
for i, t in enumerate(tasks):
    owners = [o for o in t['owners'] if o[0] != 'Unassigned']
    status = STATUSES[i % 8]
    pct = 100 if status == 'Closed' else (0 if status == 'Cancelled' else PCTS[i % 6])
    dur = (t['end'] - t['start']).days + 1 if t['start'] else None
    u = t['unit']

    def per_owner(val):
        return hm(val) if u in ('hrs/day', 'hrs') else f'{val:g}'

    def total_for(val):
        if u == 'hrs':
            return val
        perday = val if u == 'hrs/day' else val / 100 * AVAIL
        return perday * (dur or 0)

    all_o = t['owners']
    work_total = hm(sum(total_for(v) for _, _, v in all_o))
    wpo = ','.join(per_owner(v) for _, _, v in owners) if owners else (per_owner(all_o[0][2]) if all_o[0][2] else None)
    suffix = {'hrs/day': ' /day', '%/day': ' %/day', 'hrs': ''}[u]
    indiv = ','.join(f'{n}-{per_owner(v)}{suffix}' for n, _, v in owners) or None
    cs = CSETS[i % 4]
    desc = f"Workload full coverage [{t['state']}]. Scenario: {t['scenario']}. Expected: {t['expect']}."
    row = [t['name'], ', '.join(n for n, _, _ in owners) or None, status, TEAMS[i % 4], TAG[t['state']],
           ds(t['start']), ds(t['end']), str(dur) if dur else None, 'days' if dur else None,
           PRIORITIES[i % 4], str(pct), desc, work_total, wpo, u, indiv, BILLING[i % 3], t['tasklist']]
    row += list(cs) if cs else [None] * len(CUSTOM)
    ws.append(row)
    task_row.append(ws.max_row)
    for c in ws[ws.max_row]:
        c.font = Font(name=FONT)
        c.number_format = '@'

# ---------------------------------------------------------------- Read Me
rm = wb.create_sheet('Read Me')
rm.column_dimensions['A'].width = 30
rm.column_dimensions['B'].width = 120
lines = [
    ('Workload full-coverage import - Q4 2026', None),
    ('Portal / project', 'customworkload / NI-2 UIAutomation_Workload_Project2'),
    ('Built', dt.date.today().isoformat()),
    ('', ''),
    ('SHEETS', ''),
    ('Tasks', 'Copy of the import rows (the file to import is workload_full_coverage_import.xlsx, Tasks sheet only - Zoho imports rows from EVERY sheet, so never import this workbook). Same 38 columns as the original file, all text.'),
    ('Scenario Matrix', 'Planned allocation state per user, month and week, plus the edge-case list.'),
    ('Assignment Calc', 'One row per task-owner pair. Yellow cells are the settings to edit (availability, include previous import). Blue = values copied from the Tasks sheet / previous file.'),
    ('Daily Expected', 'Expected allocated hours per user per day (1 Oct - 31 Dec 2026). Red > availability, green = availability, amber under, grey weekend.'),
    ('Weekly Expected', 'Per user per Mon-Sun week: allocated hrs, available hrs, %, state.'),
    ('Monthly Expected', 'Per user per month: hrs, man-days, %, closed/cancelled hrs, day counts per state.'),
    ('', ''),
    ('COVERAGE', ''),
    ('Users', 'All 7 portal users + the resource UIAutomation_Workload_Resource + Unassigned.'),
    ('Allocation states', 'Every user gets Over, Optimal, Under and Unallocated weeks in each of Oct, Nov, Dec (rotated so users differ on the same day).'),
    ('Work units', 'hrs/day, %/day and hrs (total) - all three appear for every user across the quarter.'),
    ('Spans / edges', 'Weekend span, single day, weekend-only, zero hours, >100 %, cross-month, resource owner, multi-owner (same / different hrs, %), unassigned with and without hours, no dates.'),
    ('Field values', 'All 8 statuses, 4 priorities, 3 billing types, teams A/B/both/none, tags per state, 4 new task lists + Workload CV List, 3 custom-field value sets + blank.'),
    ('', ''),
    ('ASSUMPTIONS', ''),
    ('Availability', 'Standard Business Hours: Mon-Fri 09:00-18:00 = 9 h/day for all 7 users, no holiday list (Setup > Business Calendar, checked 08-10-2026). Resource assumed the same.'),
    ('Status handling', 'Closed and Cancelled tasks are COUNTED (observed in the report on 08-10-2026). Monthly Expected also shows the value excluding them.'),
    ('Total hours (hrs)', 'Spread evenly over the working days of the task range.'),
    ('%/day', 'Percent of the 9 h availability.'),
    ('Previous import', 'The 12 tasks already imported (UI1-T2..T13) are included in the expected values when Assignment Calc!B2 = 1. UI1-T1 and tasks in other projects are NOT included - filter the report by project NI-2 and use the WL task lists.'),
    ('VERIFY rows', 'WL-E03 (weekend-only) and WL-E10 (unassigned with hours) have uncertain product behaviour - check them rather than assert.'),
    ('', ''),
    ('IMPORT STEPS', ''),
    ('1', 'Add Vishal plreadonly, Vishal liteuser, Vishal customprofile and UIAutomation_Workload_Resource to the project first, otherwise their tasks import unassigned.'),
    ('2', 'Tasks > ... > Import Tasks > XLS/XLSX, upload workload_full_coverage_import.xlsx (NOT this workbook).'),
    ('3', 'Select date format dd/MM/yyyy.'),
    ('4', 'Statuses map 1:1. Users: choose "Map the assigned XLS Users to project users".'),
    ('5', 'Verify with Daily/Weekly/Monthly Expected.'),
]
for a, b in lines:
    rm.append([a, b])
for r in rm.iter_rows():
    for c in r:
        c.font = Font(name=FONT, bold=c.column == 1 and (c.value or '').isupper())
        c.alignment = Alignment(wrap_text=True, vertical='top')
rm['A1'].font = Font(name=FONT, bold=True, size=14)

# ---------------------------------------------------------------- Scenario Matrix
sm = wb.create_sheet('Scenario Matrix')
hdr = ['User'] + [f'{mon} W{w + 1}\n{D(2026, mm, a):%d %b}-{D(2026, mm, b):%d %b}'
                  for mon, mm, wk in MONTHS for w, (a, b) in enumerate(wk)]
sm.append(hdr)
FILL = {'Over-allocated': 'F4B6B6', 'Optimally allocated': 'B7E1B5', 'Under-allocated': 'FFE29A', 'Unallocated': 'D9D9D9'}
for _, uname in USERS:
    row = [uname]
    for mon, mm, wk in MONTHS:
        for w in range(4):
            st, v = plan[(uname, mon, w)]
            row.append(f'{st}\n{VARIANT_TEXT[st][v]}')
    sm.append(row)
    for c in sm[sm.max_row][1:]:
        st = c.value.split('\n')[0]
        c.fill = PatternFill('solid', fgColor=FILL[st])
sm.append([])
sm.append(['Edge case', 'Owners', 'Unit', 'Start', 'Due', 'Scenario', 'Expected'])
edge_hdr = sm.max_row
for t in tasks:
    if t['state'] == 'Edge':
        sm.append([t['name'], ', '.join(f'{n} ({v:g})' for n, _, v in t['owners']), t['unit'],
                   ds(t['start']), ds(t['end']), t['scenario'], t['expect']])
sm.column_dimensions['A'].width = 34
for i in range(2, 14):
    sm.column_dimensions[L(i)].width = 20
for r in sm.iter_rows():
    for c in r:
        c.font = Font(name=FONT, bold=(c.row in (1, edge_hdr)))
        c.alignment = Alignment(wrap_text=True, vertical='top')
sm.freeze_panes = 'B2'

# ---------------------------------------------------------------- Assignment Calc
ac = wb.create_sheet('Assignment Calc')
YEL = PatternFill('solid', fgColor='FFFF00')
BLUE = Font(name=FONT, color='0000FF')
ac['A1'], ac['B1'] = 'Availability (hrs/day)', AVAIL
ac['A2'], ac['B2'] = 'Include previous import (1 = yes, 0 = no)', 1
ac['C1'] = 'Source: Setup > Business Calendar > Standard Business Hours (Mon-Fri 09:00-18:00), checked 08-10-2026'
ac['C2'] = 'Previous import = the 12 tasks from the original file (UI1-T2..UI1-T13)'
for c in ('B1', 'B2'):
    ac[c].fill = YEL
    ac[c].font = BLUE
AH = ['Source', 'Tasks row', 'Task name', 'Owner (report row)', 'Status', 'Unit', 'Owner value',
      'Start', 'Due', 'Working days', 'Hrs / working day', 'Include', 'Closed/Cancelled', 'Hours in range']
ac.append([])
ac.append(AH)           # row 4
for c in ac[4]:
    c.font = bold


def calc_row(srcname, trow, name, owner, status, unit, val, start, end):
    r = ac.max_row + 1
    if trow:
        ac.append([srcname, trow, f'=Tasks!A{trow}', owner, f'=Tasks!C{trow}', f'=Tasks!O{trow}', val,
                   f'=IF(Tasks!F{trow}="","",DATE(RIGHT(Tasks!F{trow},4),MID(Tasks!F{trow},4,2),LEFT(Tasks!F{trow},2)))',
                   f'=IF(Tasks!G{trow}="","",DATE(RIGHT(Tasks!G{trow},4),MID(Tasks!G{trow},4,2),LEFT(Tasks!G{trow},2)))'])
    else:
        ac.append([srcname, None, name, owner, status, unit, val, start, end])
    ac[f'J{r}'] = f'=IF(H{r}="",0,NETWORKDAYS(H{r},I{r}))'
    ac[f'K{r}'] = (f'=IF(OR(H{r}="",J{r}=0),0,IF(F{r}="hrs/day",G{r},IF(F{r}="%/day",G{r}/100*$B$1,G{r}/J{r})))')
    ac[f'L{r}'] = f'=IF(A{r}="Previous import",$B$2,1)'
    ac[f'M{r}'] = f'=IF(OR(E{r}="Closed",E{r}="Cancelled"),1,0)'
    ac[f'N{r}'] = f'=K{r}*J{r}*L{r}'
    for c in ac[r]:
        c.font = Font(name=FONT)
    ac[f'G{r}'].font = BLUE
    if not trow:
        for col in 'CEFHI':
            ac[f'{col}{r}'].font = BLUE
    ac[f'G{r}'].comment = None
    for col in 'HI':
        ac[f'{col}{r}'].number_format = 'dd-mm-yyyy'
    ac[f'K{r}'].number_format = '0.00'
    ac[f'N{r}'].number_format = '0.00'


for t, trow in zip(tasks, task_row):
    for owner, unit, val in t['owners']:
        calc_row('New file', trow, None, owner, None, unit, val, None, None)


def to_h(s):
    h, m = s.split(':')
    return int(h) + int(m) / 60


def to_d(s):
    return dt.datetime.strptime(s, '%d/%m/%Y').date() if s else None


for r in srows[1:]:
    if not r[0] or r[0] == 'General':
        continue
    owners = [o.strip() for o in (r[1] or 'Unassigned').split(',')]
    vals = (r[13] or '0:00').split(',')
    for i, o in enumerate(owners):
        v = vals[i] if i < len(vals) else vals[0]
        val = to_h(v) if r[14] in ('hrs/day', 'hrs') else float(v)
        if r[14] == 'hrs' and r[1] is None:
            val = to_h(r[12] or '0:00')
        calc_row('Previous import', None, r[0], o, r[2], r[14], val, to_d(r[5]), to_d(r[6]))
LAST = ac.max_row
ac.column_dimensions['A'].width = 16
ac.column_dimensions['C'].width = 66
ac.column_dimensions['D'].width = 32
for col in 'BEFGHIJKLMN':
    ac.column_dimensions[col].width = 13
ac.freeze_panes = 'A5'
ac['G4'].comment = Comment('hrs/day: hours per day; %/day: percent; hrs: TOTAL hours. Copied from Tasks!N (Work Hours Per Owner) / Tasks!M.', 'Claude')


# ---------------------------------------------------------------- Daily Expected
def daily(name, closed_only):
    de = wb.create_sheet(name)
    de['A1'] = ('Expected allocated hours per day' if not closed_only else
                'Of which Closed / Cancelled hours per day') + " - formulas over 'Assignment Calc'"
    de['A1'].font = Font(name=FONT, bold=True)
    days = []
    d = D(2026, 10, 1)
    while d <= D(2026, 12, 31):
        days.append(d)
        d += dt.timedelta(days=1)
    de['A3'] = 'User / date'
    de['B3'] = 'Quarter total'
    for j, d in enumerate(days):
        c = de.cell(3, 3 + j, d)
        c.number_format = 'dd-mmm'
        de.cell(2, 3 + j, d.strftime('%a'))
    lastc = L(2 + len(days))
    rng = f'$5:${LAST}'.replace('$5', '5')
    A = "'Assignment Calc'!"
    for i, u in enumerate(REPORT_ROWS):
        r = 4 + i
        de.cell(r, 1, u)
        for j in range(len(days)):
            col = L(3 + j)
            f = (f'=IF(WEEKDAY({col}$3,2)>5,0,SUMPRODUCT(({A}$D$5:$D${LAST}=$A{r})*({A}$H$5:$H${LAST}<>"")'
                 f'*({A}$H$5:$H${LAST}<={col}$3)*({A}$I$5:$I${LAST}>={col}$3)*{A}$K$5:$K${LAST}*{A}$L$5:$L${LAST}'
                 + (f'*{A}$M$5:$M${LAST}' if closed_only else '') + '))')
            c = de.cell(r, 3 + j, f)
            c.number_format = '0.00;-0.00;-'
        de.cell(r, 2, f'=SUM(C{r}:{lastc}{r})').number_format = '0.00'
    lastr = 3 + len(REPORT_ROWS)
    for row in de.iter_rows(min_row=2, max_row=lastr):
        for c in row:
            c.font = Font(name=FONT, bold=c.row in (2, 3) or c.column == 1)
    area = f'C4:{lastc}{lastr}'
    de.conditional_formatting.add(area, FormulaRule(formula=['WEEKDAY(C$3,2)>5'], fill=PatternFill('solid', fgColor='D9D9D9'), stopIfTrue=True))
    if not closed_only:
        de.conditional_formatting.add(area, FormulaRule(formula=["AND($A4<>\"Unassigned\",C4>'Assignment Calc'!$B$1+0.0001)"], fill=PatternFill('solid', fgColor='F4B6B6'), stopIfTrue=True))
        de.conditional_formatting.add(area, FormulaRule(formula=["AND($A4<>\"Unassigned\",ABS(C4-'Assignment Calc'!$B$1)<0.0001)"], fill=PatternFill('solid', fgColor='B7E1B5'), stopIfTrue=True))
        de.conditional_formatting.add(area, FormulaRule(formula=['C4>0'], fill=PatternFill('solid', fgColor='FFE29A'), stopIfTrue=True))
    de.column_dimensions['A'].width = 32
    de.column_dimensions['B'].width = 12
    for j in range(len(days)):
        de.column_dimensions[L(3 + j)].width = 7
    de.freeze_panes = 'C4'
    return de, days, lastc, lastr


de, days, DLC, DLR = daily('Daily Expected', False)
dc, *_ = daily('Daily Closed-Cancelled', True)


# ---------------------------------------------------------------- Weekly Expected
we = wb.create_sheet('Weekly Expected')
we.append(['User', 'Week start (Mon)', 'Week end (Sun)', 'Working days in Q4', 'Allocated hrs',
           'Available hrs', 'Allocation %', 'State', 'Closed/Cancelled hrs'])
for c in we[1]:
    c.font = bold
DA = f"'Daily Expected'!"
weeks = []
d = D(2026, 9, 28)
while d <= D(2026, 12, 28):
    weeks.append(d)
    d += dt.timedelta(days=7)
for u in REPORT_ROWS:
    for wk in weeks:
        r = we.max_row + 1
        we.append([u, wk, wk + dt.timedelta(days=6)])
        we[f'D{r}'] = f'=NETWORKDAYS(MAX(B{r},DATE(2026,10,1)),MIN(C{r},DATE(2026,12,31)))'
        sp = (f"(({{S}}$A$4:$A${DLR}=$A{r})*({{S}}$C$3:${DLC}$3>=$B{r})*({{S}}$C$3:${DLC}$3<=$C{r})*{{S}}$C$4:${DLC}${DLR})")
        we[f'E{r}'] = '=SUMPRODUCT' + sp.replace('{S}', DA)
        we[f'F{r}'] = f"=IF($A{r}=\"Unassigned\",0,D{r}*'Assignment Calc'!$B$1)"
        we[f'G{r}'] = f'=IF(F{r}=0,0,E{r}/F{r})'
        we[f'H{r}'] = (f'=IF($A{r}="Unassigned",IF(E{r}>0,"Unassigned hours","-"),IF(E{r}=0,"Unallocated",'
                       f'IF(E{r}>F{r}+0.0001,"Over-allocated",IF(ABS(E{r}-F{r})<0.0001,"Optimally allocated","Under-allocated"))))')
        we[f'I{r}'] = '=SUMPRODUCT' + sp.replace('{S}', "'Daily Closed-Cancelled'!")
        for c in we[r]:
            c.font = Font(name=FONT)
        for col in 'BC':
            we[f'{col}{r}'].number_format = 'dd-mm-yyyy'
        for col in 'EFI':
            we[f'{col}{r}'].number_format = '0.00'
        we[f'G{r}'].number_format = '0.00%'
WLR = we.max_row
for st, color in FILL.items():
    we.conditional_formatting.add(f'H2:H{WLR}', FormulaRule(formula=[f'H2="{st}"'], fill=PatternFill('solid', fgColor=color)))
note = WLR + 2
we[f'A{note}'] = ('Note: a week is mixed when daily states differ (e.g. an edge task on 1-2 Oct). '
                  'State uses the weekly total vs weekly availability, which is how the Weeks view buckets data.')
we[f'A{note}'].font = Font(name=FONT, italic=True)
for col, w in zip('ABCDEFGHI', [32, 15, 15, 12, 13, 13, 12, 20, 15]):
    we.column_dimensions[col].width = w
we.freeze_panes = 'B2'
we.auto_filter.ref = f'A1:I{WLR}'

# ---------------------------------------------------------------- Monthly Expected
me = wb.create_sheet('Monthly Expected')
me.append(['User', 'Month', 'From', 'To', 'Working days', 'Allocated hrs', 'Allocated man-days',
           'Available hrs', 'Allocation %', 'State (month total)', 'Closed/Cancelled hrs', 'Allocated excl. Closed/Cancelled',
           'Over days', 'Optimal days', 'Under days', 'Unallocated working days'])
for c in me[1]:
    c.font = bold
periods = [('Oct 2026', D(2026, 10, 1), D(2026, 10, 31)), ('Nov 2026', D(2026, 11, 1), D(2026, 11, 30)),
           ('Dec 2026', D(2026, 12, 1), D(2026, 12, 31)), ('Q4 2026', D(2026, 10, 1), D(2026, 12, 31))]
AV = "'Assignment Calc'!$B$1"
for u in REPORT_ROWS:
    for pname, a, b in periods:
        r = me.max_row + 1
        me.append([u, pname, a, b])
        me[f'E{r}'] = f'=NETWORKDAYS(C{r},D{r})'
        base = (f"({DA}$A$4:$A${DLR}=$A{r})*({DA}$C$3:${DLC}$3>=$C{r})*({DA}$C$3:${DLC}$3<=$D{r})")
        wkday = f"(WEEKDAY({DA}$C$3:${DLC}$3,2)<6)"
        vals = f"{DA}$C$4:${DLC}${DLR}"
        me[f'F{r}'] = f'=SUMPRODUCT({base}*{vals})'
        me[f'G{r}'] = f'=F{r}/{AV}'
        me[f'H{r}'] = f'=IF($A{r}="Unassigned",0,E{r}*{AV})'
        me[f'I{r}'] = f'=IF(H{r}=0,0,F{r}/H{r})'
        me[f'J{r}'] = (f'=IF($A{r}="Unassigned",IF(F{r}>0,"Unassigned hours","-"),IF(F{r}=0,"Unallocated",'
                       f'IF(F{r}>H{r}+0.0001,"Over-allocated",IF(ABS(F{r}-H{r})<0.0001,"Optimally allocated","Under-allocated"))))')
        me[f'K{r}'] = (f"=SUMPRODUCT(('Daily Closed-Cancelled'!$A$4:$A${DLR}=$A{r})*('Daily Closed-Cancelled'!$C$3:${DLC}$3>=$C{r})"
                       f"*('Daily Closed-Cancelled'!$C$3:${DLC}$3<=$D{r})*'Daily Closed-Cancelled'!$C$4:${DLC}${DLR})")
        me[f'L{r}'] = f'=F{r}-K{r}'
        me[f'M{r}'] = f'=SUMPRODUCT({base}*{wkday}*({vals}>{AV}+0.0001))'
        me[f'N{r}'] = f'=SUMPRODUCT({base}*{wkday}*(ABS({vals}-{AV})<0.0001))'
        me[f'O{r}'] = f'=SUMPRODUCT({base}*{wkday}*({vals}>0)*({vals}<{AV}-0.0001))'
        me[f'P{r}'] = f'=SUMPRODUCT({base}*{wkday}*({vals}=0))'
        for c in me[r]:
            c.font = Font(name=FONT, bold=pname.startswith('Q4'))
        for col in 'CD':
            me[f'{col}{r}'].number_format = 'dd-mm-yyyy'
        for col in 'FGHKL':
            me[f'{col}{r}'].number_format = '0.00'
        me[f'I{r}'].number_format = '0.00%'
MLR = me.max_row
for st, color in FILL.items():
    me.conditional_formatting.add(f'J2:J{MLR}', FormulaRule(formula=[f'J2="{st}"'], fill=PatternFill('solid', fgColor=color)))
for col, w in zip('ABCDEFGHIJKLMNOP', [32, 10, 12, 12, 9, 12, 12, 12, 11, 20, 13, 14, 9, 9, 9, 12]):
    me.column_dimensions[col].width = w
me.freeze_panes = 'C2'
me.auto_filter.ref = f'A1:P{MLR}'

wb.save(OUT)
imp = openpyxl.Workbook()
iw = imp.active
iw.title = 'Tasks'
for row in ws.iter_rows(values_only=True):
    iw.append(list(row))
for c in iw[1]:
    c.font = bold
for r in iw.iter_rows(min_row=2):
    for c in r:
        c.font = Font(name=FONT)
        c.number_format = '@'
for col, dim in ws.column_dimensions.items():
    iw.column_dimensions[col].width = dim.width
iw.freeze_panes = 'B2'
imp.save(OUT_IMPORT)
print('tasks', len(tasks), 'calc rows', LAST - 4, 'out', OUT)
