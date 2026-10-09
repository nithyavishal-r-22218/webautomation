"""Run report for the API-driven workload run (project NI-8, 08-10-2026)."""
import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule

c = json.load(open('cases.json'))
OUT = '/mnt/user-data/outputs/workload_api_run_results.xlsx'
F = 'Arial'
wl_fail = {
    'Vishal PLAdmin Oct W4 Over-allocated': ('F1', '5/5 days: 540 min vs 810 expected'),
    'vishal PL manager Nov W2 Over-allocated': ('F1', '5/5 days: 540 vs 810'),
    'Vishal plemployee Dec W4 Over-allocated': ('F1', '4/4 days: 540 vs 810'),
    'Vishal liteuser Oct W4 Over-allocated': ('F1', '5/5 days: 540 vs 810'),
    'Vishal customprofile Nov W2 Over-allocated': ('F1', '5/5 days: 540 vs 810'),
    'UIAutomation_Workload_Resource Dec W4 Over-allocated': ('F1', '4/4 days: 540 vs 810'),
    'WL-E03 Weekend-only task - vishal PL manager': ('F2', 'Sat 03-10 and Sun 04-10: 240 min each vs 0 expected'),
    'WL-E05 150 percent per day - Vishal customprofile': ('F1', '01-10 and 02-10: 540 vs 810'),
}
field_fail = {
    'WL Nov W4 Over-allocated (1 of 2) - Vishal workload': 'pct', 'WL Oct W4 Over-allocated - Vishal PLAdmin': 'work+total+pct',
    'WL Dec W3 Optimally allocated - Vishal PLAdmin': 'pct', 'WL Nov W2 Over-allocated - vishal PL manager': 'work+total',
    'WL Nov W4 Under-allocated - vishal PL manager': 'pct', 'WL Oct W3 Optimally allocated - Vishal plemployee': 'pct',
    'WL Dec W4 Over-allocated - Vishal plemployee': 'work+total+pct', 'WL Nov W4 Over-allocated (1 of 2) - Vishal plreadonly': 'pct',
    'WL Oct W4 Over-allocated - Vishal liteuser': 'work+total+pct', 'WL Dec W3 Optimally allocated - Vishal liteuser': 'pct',
    'WL Nov W2 Over-allocated - Vishal customprofile': 'work+total', 'WL Nov W4 Under-allocated - Vishal customprofile': 'pct',
    'WL Oct W3 Optimally allocated - UIAutomation_Workload_Resource': 'pct', 'WL Dec W4 Over-allocated - UIAutomation_Workload_Resource': 'work+total+pct',
    'WL-E05 150 percent per day - Vishal customprofile': 'work+total', 'WL-E08 Multi-owner different hours - Vishal liteuser and Vishal customprofile': 'pct',
}
tot_fail = {('Vishal PLAdmin', 'Oct'): (-22.5, 'F1'), ('vishal PL manager', 'Oct'): (8, 'F2'), ('vishal PL manager', 'Nov'): (-22.5, 'F1'),
            ('Vishal plemployee', 'Dec'): (-18, 'F1'), ('Vishal liteuser', 'Oct'): (-22.5, 'F1'), ('Vishal customprofile', 'Oct'): (-9, 'F1'),
            ('Vishal customprofile', 'Nov'): (-22.5, 'F1'), ('UIAutomation_Workload_Resource', 'Dec'): (-18, 'F1')}

wb = openpyxl.Workbook()
s = wb.active
s.title = 'Summary'
rows = [
    ['Workload API run - project NI-8 WL_API_FullCoverage_20261008112546 (id 306499000000174009), portal customworkload, 08-10-2026'],
    [],
    ['Step', 'Call / action', 'Result'],
    ['1 Create project', 'POST /api/v3/portal/131176135/projects (JSON: name, description, added_via=api, start 2026-10-01, end 2026-12-31)', '201 - NI-8 created; GET 200, owner Vishal workload, Standard Business Hours'],
    ['2 Associate users', 'POST /projects/{id}/projectusers (form-encoded userdetails JSON array of email_id/profile_id/role_id, notify=false)', '201 - 6 users added (PLAdmin, plemployee, PL manager, plreadonly, customprofile, liteuser); GET /projects/{id}/users = 7 users with their own profiles. JSON body was rejected; form-encoded worked.'],
    ['2b Resource', 'UI Add Resource (app calls POST /projects/{id}/resources - not in the public docs)', '204 - UIAutomation_Workload_Resource associated'],
    ['3 Import tasks', 'UI import: workload_full_coverage_import.xlsx, dd/MM/yyyy, XLS users mapped to project users', 'Import Successful; GET /projects/{id}/tasks = 111 tasks'],
    ['4 Execute cases', 'GET /projects/{id}/tasks (field cases) + GET /reports/workload?project_id={id}&time_period=days for Oct, Nov, Dec (+ fetch_unassignee=true)', 'See counts below'],
    ['5 Remove project', 'POST /projects/{id}/trash', '204; GET /projects/{id} now 410 RESOURCE_TRASHED; not in active list. Permanent DELETE /projects/{id} not executed (irreversible) - restore possible via POST /projects/{id}/restore'],
    [],
    ['Case group', 'Total', 'Pass', 'Fail', 'Fail explained by'],
    ['Task field cases (owners, unit, work per owner, total work, dates, status, task list, priority, % complete, billing)', 111, 111 - len(field_fail), len(field_fail), 'F1 x7, F6 x13 (some rows have both)'],
    ['Workload day cases (per user-week scenario + edge cases + unallocated weeks; hours AND over/optimal/under label per day)', 111, 111 - len(wl_fail), len(wl_fail), 'F1 x7, F2 x1'],
    ['Monthly totals per user (8 owners x 3 months + Unassigned)', 27, 27 - len(tot_fail), len(tot_fail), 'F1 x7, F2 x1'],
    ['Unexpected allocation days (any non-zero day not planned)', '-', '0 found', '-', 'No-date task WL-E12 correctly absent'],
    [],
    ['FINDINGS'],
    ['F1', 'Import caps Work %/day at 100. Every 150 %/day row stored as 100 % (work_values 100, total 45:00 instead of 67:30); Workload shows 9:00/day (optimal) instead of 13:30 (over). No warning in the import wizard.'],
    ['F2', 'Task whose whole range is a weekend (WL-E03, Sat 03-10 to Sun 04-10, 4:00 hrs/day) is allocated 4:00 on Saturday and Sunday in the Workload report although business hours are Mon-Fri.'],
    ['F6', 'Cancelled tasks are stored with 100 % completion although the file says 0 % (13 tasks). Likely by design (Cancelled is a closed-type status) - confirm.'],
    ['F7', 'POST /projectusers rejects a JSON body; works only form-encoded (the docs request example uses --form while setting Content-Type: application/json).'],
    ['F8', 'No public API to associate a resource with a project; the UI uses POST /projects/{id}/resources.'],
    ['F9', 'Trashed project still listed in the sidebar Recent Projects.'],
    ['OK', 'Lite user, read-only user, custom-profile user and resource can own tasks; hrs/day, total hrs and %/day all allocate correctly up to 100 %; multi-owner split (WL-E08, E09, E15); cross-month (E06, E07); weekend span (E01); unassigned with hours (E10 = 4:00/day under Unassigned); labels over/optimal/under correct on every passing day.'],
    ['Not checked here', 'Associated Team and Tags are not returned by GET /tasks; on the NI-2 import they were empty in the UI (earlier finding F3).'],
]
for r in rows:
    s.append(r)
for r in s.iter_rows():
    for x in r:
        x.font = Font(name=F, bold=(x.row in (1, 3, 11, 17)) or (x.column == 1 and x.row > 17))
        x.alignment = Alignment(wrap_text=True, vertical='top')
for col, w in zip('ABCDE', [34, 70, 70, 10, 34]):
    s.column_dimensions[col].width = w
s['A1'].font = Font(name=F, bold=True, size=13)

w = wb.create_sheet('Workload cases')
w.append(['Case', 'Type', 'Tasks', 'Result', 'Known finding', 'Detail'])
for k in c['cases']:
    f = wl_fail.get(k['id'])
    w.append([k['id'], k['kind'], '; '.join(k['tasks']) or '(no task - expect 0)', 'FAIL' if f else 'PASS', f[0] if f else '', f[1] if f else f"{len(k['checks'])} day checks matched (hours + label)"])
fl = wb.create_sheet('Field cases')
fl.append(['Task', 'Result', 'Failed fields', 'Known finding'])
for t in c['tasks']:
    ff = field_fail.get(t['name'])
    kn = []
    if ff and ('work' in ff):
        kn.append('F1')
    if ff and 'pct' in ff:
        kn.append('F6')
    fl.append([t['name'], 'FAIL' if ff else 'PASS', ff or '', ','.join(kn)])
mt = wb.create_sheet('Monthly totals')
mt.append(['User', 'Month', 'Expected hrs', 'Actual hrs', 'Diff hrs', 'Result', 'Known finding'])
for t in c['totals']:
    d = tot_fail.get((t['u'], t['mon']))
    exp = t['m'] / 60
    mt.append([t['u'], t['mon'], round(exp, 2), round(exp + (d[0] if d else 0), 2), d[0] if d else 0, 'FAIL' if d else 'PASS', d[1] if d else ''])
for sh, col in ((w, 'D'), (fl, 'B'), (mt, 'F')):
    for x in sh[1]:
        x.font = Font(name=F, bold=True)
    for r in sh.iter_rows(min_row=2):
        for x in r:
            x.font = Font(name=F)
    n = sh.max_row
    sh.conditional_formatting.add(f'{col}2:{col}{n}', FormulaRule(formula=[f'{col}2="PASS"'], fill=PatternFill('solid', fgColor='B7E1B5')))
    sh.conditional_formatting.add(f'{col}2:{col}{n}', FormulaRule(formula=[f'{col}2="FAIL"'], fill=PatternFill('solid', fgColor='F4B6B6')))
    sh.freeze_panes = 'B2'
    sh.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(sh.max_column)}{n}'
w.column_dimensions['A'].width = 58; w.column_dimensions['C'].width = 80; w.column_dimensions['F'].width = 50
fl.column_dimensions['A'].width = 80; fl.column_dimensions['C'].width = 22
mt.column_dimensions['A'].width = 32
wb.save(OUT)
print('saved', OUT)
